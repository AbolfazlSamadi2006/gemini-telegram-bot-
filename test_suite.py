"""
Exhaustive 5-Tier Adversarial Stress-Testing and QA Test Suite (test_suite.py)
Requirement R7 - Milestone 5 Verification Suite

Architecture & Test Philosophy:
- 100% Offline & Deterministic: High-fidelity mocks for Google GenAI SDK & Telegram Bot API.
- Zero reliance on paid API credentials or live external network connections.
- 5 Distinct Test Tiers:
    Tier 1: Feature Isolation & Unit Verification (R1, R2, R3, R4, R5, R6)
    Tier 2: Boundary & Corner Cases (20MB files, corrupt payloads, 404/500 URLs, midnight clock, intent isolation)
    Tier 3: Concurrency & Stress Testing (50 concurrent SQLite writers, zero database locks, WAL mode)
    Tier 4: Resilience & Fault Tolerance (Gemini 503 backoff, 429 rate limit backoff, server reboot recovery, Markdown parse error fallback)
    Tier 5: Real-World Workload Scenarios (End-to-end multi-turn conversation exercising math, charts, exports, reminders, and reasoning)

Execution:
    python test_suite.py
    python -m unittest test_suite.py
"""

import os
import sys
import io
import time
import uuid
import re
import asyncio
import sqlite3
import datetime
import unittest
import threading
from typing import List, Dict, Optional, Tuple, Any
from unittest.mock import MagicMock, AsyncMock, patch

# Ensure UTF-8 output encoding on Windows consoles
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

# Force offline environment variables
os.environ.setdefault("GEMINI_API_KEY", "dummy_qa_suite_gemini_key_offline")
os.environ.setdefault("TELEGRAM_BOT_TOKEN", "123456789:dummy_qa_suite_token_offline")
os.environ.setdefault("ALLOWED_USERS", "")

# Import target application under test
import bot
from bot import (
    Database,
    ChatResponse,
    TEHRAN_TZ,
    normalize_digits,
    has_reminder_intent,
    clean_reminder_text,
    parse_reminder_input_tier1,
    parse_reminder_input,
    _execute_gemini_reminder_fallback,
    parse_gemini_candidate_parts,
    format_response_with_code_execution,
    build_gemini_session,
    generate_chat_response,
    generate_multimodal_response,
    handle_draw_command,
    enhance_image_prompt,
    extract_urls,
    extract_youtube_video_id,
    fetch_youtube_transcript,
    fetch_web_page_content,
    fetch_and_summarize_url,
    get_persian_font_path,
    setup_pdf_fonts,
    reshape_persian_text,
    generate_excel_report,
    generate_pdf_document,
    generate_chart,
    extract_markdown_table,
    detect_export_intent,
    perform_chat_export,
    export_chat,
    handle_message,
    handle_document,
    handle_audio,
    handle_draw_command,
    handle_remind_command,
    callback_handler,
    reminder_scheduler_loop,
    recover_server_reboot_reminders,
    extract_docx_text,
    AVAILABLE_MODELS,
    FALLBACK_MODELS,
)

import openpyxl
import httpx
from google.genai import types
from telegram import Update, Message, User, Chat, CallbackQuery, InlineKeyboardMarkup, InlineKeyboardButton
from telegram.error import BadRequest, Forbidden


# =============================================================================
# High-Fidelity Test Doubles & Mocks
# =============================================================================

class MockGenAIPart:
    """High-fidelity double for google.genai.types.Part."""
    def __init__(
        self,
        text: Optional[str] = None,
        thought: bool = False,
        executable_code: Optional[str] = None,
        code_execution_result: Optional[str] = None,
    ):
        self.text = text
        self.thought = thought
        if executable_code is not None:
            self.executable_code = MagicMock()
            self.executable_code.code = executable_code
        else:
            self.executable_code = None

        if code_execution_result is not None:
            self.code_execution_result = MagicMock()
            self.code_execution_result.output = code_execution_result
        else:
            self.code_execution_result = None


class MockGenAIResponse:
    """High-fidelity double for Gemini API GenerateContentResponse."""
    def __init__(self, parts: List[MockGenAIPart], text: str = ""):
        self.text = text
        mock_content = MagicMock()
        mock_content.parts = parts
        mock_candidate = MagicMock()
        mock_candidate.content = mock_content
        self.candidates = [mock_candidate]


class MockTelegramFile:
    """High-fidelity double for Telegram File download."""
    def __init__(self, data: bytes):
        self._data = data

    async def download_as_bytearray(self) -> bytearray:
        return bytearray(self._data)


class MockTelegramDocument:
    """High-fidelity double for telegram.Document."""
    def __init__(self, file_name: str, file_size: int, data: bytes = b""):
        self.file_name = file_name
        self.file_size = file_size
        self.mime_type = "application/octet-stream"
        self._data = data

    async def get_file(self) -> MockTelegramFile:
        return MockTelegramFile(self._data)


class MockTelegramAudio:
    """High-fidelity double for telegram.Audio or telegram.Voice."""
    def __init__(self, file_name: str, file_size: int, is_voice: bool = False, data: bytes = b""):
        self.file_name = file_name
        self.file_size = file_size
        self._data = data
        self.is_voice = is_voice

    async def get_file(self) -> MockTelegramFile:
        return MockTelegramFile(self._data)


def create_mock_update_message(
    user_id: int = 12345678,
    chat_id: int = 98765432,
    text: Optional[str] = None,
    document: Optional[MockTelegramDocument] = None,
    audio: Optional[MockTelegramAudio] = None,
    voice: Optional[MockTelegramAudio] = None,
    message_id: int = 100,
) -> Tuple[MagicMock, MagicMock]:
    """Helper to create realistic Mock Update and Context pairs."""
    mock_update = MagicMock(spec=Update)
    mock_update.effective_user = MagicMock(spec=User, id=user_id)
    mock_update.effective_chat = MagicMock(spec=Chat, id=chat_id)

    mock_msg = MagicMock(spec=Message)
    mock_msg.message_id = message_id
    mock_msg.text = text
    mock_msg.caption = None
    mock_msg.document = document
    mock_msg.audio = audio
    mock_msg.voice = voice

    # Track responses
    mock_status_msg = MagicMock()
    mock_status_msg.edit_text = AsyncMock()
    mock_status_msg.delete = AsyncMock()

    mock_msg.reply_text = AsyncMock(return_value=mock_status_msg)
    mock_msg.reply_photo = AsyncMock()
    mock_msg.reply_document = AsyncMock()

    mock_update.message = mock_msg

    mock_context = MagicMock()
    mock_context.bot = MagicMock()
    mock_context.bot.send_message = AsyncMock(return_value=mock_msg)
    mock_context.bot.send_chat_action = AsyncMock()
    mock_context.bot.send_photo = AsyncMock()
    mock_context.bot.send_document = AsyncMock()

    return mock_update, mock_context


# =============================================================================
# Tier 1: Feature Isolation & Unit Verification
# =============================================================================

class Tier1FeatureIsolationTests(unittest.IsolatedAsyncioTestCase):
    """
    Tier 1: Feature Isolation & Unit Verification
    Verifies R1, R2, R3, R4, R5, and R6 in strict isolation with zero network access.
    """
    tier = 1

    def setUp(self):
        self.test_db_path = f"test_tier1_{uuid.uuid4().hex[:8]}.db"
        self.db = Database(self.test_db_path)
        bot.db = self.db
        self.user_id = "11001"
        self.db.get_or_create_user(self.user_id)
        self.active_chat = self.db.get_active_chat(self.user_id)

    def tearDown(self):
        for extra in ["", "-wal", "-shm"]:
            p = self.test_db_path + extra
            if os.path.exists(p):
                try:
                    os.remove(p)
                except Exception:
                    pass

    async def test_t1_01_code_execution_badge_parsing(self):
        """[R1] Verify Python code execution tool parsing, execution result, and Persian badges."""
        code_snip = "import math\nresult = math.factorial(6)\nprint(result)"
        exec_out = "720\n"
        clean_ans = "فاکتوریل عدد ۶ برابر ۷۲۰ است."
        thought = "محاسبه دقیق مقدار فاکتوریل با کد پایتون بدون خطا."

        mock_parts = [
            MockGenAIPart(thought=True, text=thought),
            MockGenAIPart(executable_code=code_snip),
            MockGenAIPart(code_execution_result=exec_out),
            MockGenAIPart(text=clean_ans),
        ]
        mock_resp = MockGenAIResponse(parts=mock_parts, text=clean_ans)

        chat_resp = parse_gemini_candidate_parts(mock_resp)
        self.assertEqual(chat_resp.clean_text, clean_ans)
        self.assertEqual(chat_resp.executed_code, code_snip)
        self.assertEqual(chat_resp.execution_result, exec_out.strip())
        self.assertEqual(chat_resp.thought_text, thought)

        # Verify Persian Badge Output formatting
        formatted = chat_resp.formatted_text
        self.assertIn("💻 **کد پایتون اجرا شده:**", formatted)
        self.assertIn("📊 **نتیجه تایید شده اجرا:**", formatted)
        self.assertIn("math.factorial(6)", formatted)
        self.assertIn("720", formatted)

    async def test_t1_02_imagen3_aspect_ratios_and_prompt_enhancement(self):
        """[R2] Verify Imagen 3 /draw and /image command with aspect ratio flags and photo delivery."""
        mock_update, mock_context = create_mock_update_message(
            text="/draw غروب آفتاب در سواحل خلیج فارس --ar 16:9"
        )

        fake_jpeg = b"\xFF\xD8\xFF\xE0\x00\x10JFIF\x00\x01\x01\x01\x00`\x00`\x00\x00\xFF\xDB"
        mock_gen_img = MagicMock()
        mock_gen_img.rai_filtered_reason = None
        mock_gen_img.image.image_bytes = fake_jpeg
        mock_gen_response = MagicMock(generated_images=[mock_gen_img])

        with patch.object(bot.client.models, "generate_images", return_value=mock_gen_response) as mock_img_gen, \
             patch("bot.enhance_image_prompt", AsyncMock(return_value="A breathtaking cinematic sunset over Persian Gulf, 8k")):

            await handle_draw_command(mock_update, mock_context)

            # Verify client.models.generate_images call
            mock_img_gen.assert_called_once()
            _, kwargs = mock_img_gen.call_args
            self.assertEqual(kwargs["model"], "imagen-3.0-generate-002")
            self.assertEqual(kwargs["config"].aspect_ratio, "16:9")
            self.assertEqual(kwargs["config"].number_of_images, 1)

            # Verify Telegram photo delivery
            mock_update.message.reply_photo.assert_called_once()
            _, photo_kwargs = mock_update.message.reply_photo.call_args
            self.assertIn("16:9", photo_kwargs["caption"])
            self.assertIn("Imagen 3", photo_kwargs["caption"])
            photo_file = photo_kwargs["photo"]
            self.assertEqual(photo_file.getvalue(), fake_jpeg)

    async def test_t1_03_bilingual_reminder_parser_persian_and_english(self):
        """[R3] Verify bilingual reminder parser supporting Persian/ASCII digits, relative delays, and clock times."""
        base_dt = datetime.datetime(2026, 10, 16, 12, 0, 0, tzinfo=TEHRAN_TZ)  # 12:00 Tehran = 08:30 UTC

        # Case A: Persian digits and relative minutes
        res_a = parse_reminder_input_tier1("۱۰ دقیقه دیگه تماس با مشتری", base_dt)
        self.assertIsNotNone(res_a)
        due_utc_a, text_a = res_a
        self.assertEqual(due_utc_a, "2026-10-16T08:40:00Z")
        self.assertEqual(text_a, "تماس با مشتری")

        # Case B: English command shorthand with hours
        res_b = parse_reminder_input_tier1("/remind 2h review pull request", base_dt)
        self.assertIsNotNone(res_b)
        due_utc_b, text_b = res_b
        self.assertEqual(due_utc_b, "2026-10-16T10:30:00Z")
        self.assertEqual(text_b, "review pull request")

        # Case C: Persian clock time with tomorrow marker
        res_c = parse_reminder_input_tier1("فردا ساعت ۹ صبح آزمون تافل", base_dt)
        self.assertIsNotNone(res_c)
        due_utc_c, text_c = res_c
        self.assertEqual(due_utc_c, "2026-10-17T05:30:00Z")
        self.assertEqual(text_c, "آزمون تافل")

    async def test_t1_04_tehran_timezone_utc_conversion(self):
        """[R3] Verify strict Tehran UTC+03:30 timezone offset and ISO-8601 UTC formatting."""
        # Standard Tehran timezone offset must be +03:30
        offset = TEHRAN_TZ.utcoffset(None)
        self.assertEqual(offset, datetime.timedelta(hours=3, minutes=30))

        # Local time 14:00 Tehran must map to 10:30 UTC
        local_tehran = datetime.datetime(2026, 10, 16, 14, 0, 0, tzinfo=TEHRAN_TZ)
        utc_dt = local_tehran.astimezone(datetime.timezone.utc)
        self.assertEqual(utc_dt.hour, 10)
        self.assertEqual(utc_dt.minute, 30)
        self.assertEqual(utc_dt.strftime("%Y-%m-%dT%H:%M:%SZ"), "2026-10-16T10:30:00Z")

    async def test_t1_05_url_and_youtube_extraction(self):
        """[R4] Verify URL regex detection, YouTube video ID extraction, and HTML article parsing."""
        # Multi-URL extraction
        raw_text = "لطفاً این دو لینک را مطالعه کن: https://news.ycombinator.com/item?id=1234 و https://youtu.be/dQw4w9WgXcQ?t=10 ممنون!"
        urls = extract_urls(raw_text)
        self.assertEqual(len(urls), 2)
        self.assertEqual(urls[0], "https://news.ycombinator.com/item?id=1234")
        self.assertEqual(urls[1], "https://youtu.be/dQw4w9WgXcQ?t=10")

        # YouTube video ID extraction
        yt_id = extract_youtube_video_id(urls[1])
        self.assertEqual(yt_id, "dQw4w9WgXcQ")

        # HTML extraction via simulated web page
        html_doc = """
        <html>
            <head><title>هوش مصنوعی در پزشکی</title><meta name="description" content="کاربردهای نوین"></head>
            <body>
                <header><nav>منو</nav></header>
                <article>
                    <h1>تحول هوش مصنوعی</h1>
                    <p>هوش مصنوعی تشخیص بیماری‌ها را با دقت بالای ۹۹٪ بهبود بخشیده است.</p>
                </article>
                <footer>کپی رایت ۲۰۲۶</footer>
            </body>
        </html>
        """
        mock_http_resp = MagicMock(status_code=200, text=html_doc, headers={"content-type": "text/html; charset=utf-8"})
        with patch("httpx.AsyncClient.get", AsyncMock(return_value=mock_http_resp)):
            page_info, err = await fetch_web_page_content("https://med-ai.org/article")
            self.assertIsNone(err)
            self.assertIsNotNone(page_info)
            self.assertEqual(page_info["title"], "هوش مصنوعی در پزشکی")
            self.assertEqual(page_info["meta_description"], "کاربردهای نوین")
            self.assertIn("هوش مصنوعی تشخیص بیماری‌ها", page_info["text"])
            self.assertNotIn("کپی رایت", page_info["text"])

    async def test_t1_06_dynamic_excel_pdf_chart_generation(self):
        """[R5] Verify dynamic file generation: Excel (.xlsx with RTL), PDF (%PDF- header), and PNG chart."""
        # 1. Excel Generation
        dataset = [
            {"ردیف": 1, "شاخص": "دقت مدل", "مقدار": "99.4%"},
            {"ردیف": 2, "شاخص": "سرعت پاسخ", "مقدار": "120ms"},
        ]
        excel_buf = generate_excel_report(dataset, title="گزارش_کیفی")
        self.assertIsInstance(excel_buf, io.BytesIO)
        self.assertTrue(excel_buf.name.endswith(".xlsx"))
        self.assertTrue(excel_buf.getvalue().startswith(b"PK\x03\x04"))

        wb = openpyxl.load_workbook(excel_buf)
        ws = wb.active
        self.assertTrue(ws.sheet_view.rightToLeft)
        self.assertEqual(ws.max_row, 3)

        # 2. PDF Generation with Persian RTL text
        pdf_buf = generate_pdf_document("گزارش تحلیلی", "این یک گزارش رسمی جهت تست ماژول پی‌دی‌اف است.")
        self.assertIsInstance(pdf_buf, io.BytesIO)
        self.assertTrue(pdf_buf.name.endswith(".pdf"))
        self.assertTrue(pdf_buf.getvalue().startswith(b"%PDF-"))

        # 3. Chart Generation (Bar, Line, Pie)
        chart_data = {"labels": ["بهار", "تابستان", "پاییز"], "values": [120, 250, 180]}
        for c_type in ["bar", "line", "pie"]:
            c_buf = generate_chart(chart_data, chart_type=c_type)
            self.assertIsInstance(c_buf, io.BytesIO)
            self.assertTrue(c_buf.getvalue().startswith(b"\x89PNG\r\n\x1a\n"))

    async def test_t1_07_transparent_reasoning_logging_and_display(self):
        """[R6] Verify ThinkingConfig logging in SQLite reasoning_logs and inline button inspection."""
        thought_id = uuid.uuid4().hex[:12]
        thought_content = "ابتدا مساله را تحلیل می‌کنم. گام اول: استخراج پارامترها. گام دوم: حل معادله دیفرانسیل."

        # Save thought
        self.db.save_thought(thought_id, self.user_id, self.active_chat["id"], thought_content)

        # Retrieve thought from DB
        fetched_thought = self.db.get_thought(thought_id)
        self.assertEqual(fetched_thought, thought_content)

        # Verify Telegram Callback Handler retrieves and displays the thought
        mock_query = MagicMock(spec=CallbackQuery)
        mock_query.data = f"show_thought:{thought_id}"
        mock_query.message = MagicMock(spec=Message)
        mock_query.message.reply_text = AsyncMock()
        mock_query.answer = AsyncMock()

        mock_update = MagicMock(spec=Update)
        mock_update.callback_query = mock_query
        mock_context = MagicMock()

        await callback_handler(mock_update, mock_context)

        mock_query.answer.assert_called_once()
        mock_query.message.reply_text.assert_called_once()
        reply_call_args = mock_query.message.reply_text.call_args[0][0]
        self.assertIn("استخراج پارامترها", reply_call_args)
        self.assertIn("فرآیند و زنجیره تفکر", reply_call_args)


# =============================================================================
# Tier 2: Boundary & Corner Cases
# =============================================================================

class Tier2BoundaryCornerCaseTests(unittest.IsolatedAsyncioTestCase):
    """
    Tier 2: Boundary & Corner Cases
    Verifies 20MB payload limit, corrupted files, 404/500 URLs, midnight clock, and intent isolation.
    """
    tier = 2

    def setUp(self):
        self.test_db_path = f"test_tier2_{uuid.uuid4().hex[:8]}.db"
        self.db = Database(self.test_db_path)
        bot.db = self.db
        self.user_id = "12002"
        self.db.get_or_create_user(self.user_id)
        self.active_chat = self.db.get_active_chat(self.user_id)

    def tearDown(self):
        for extra in ["", "-wal", "-shm"]:
            p = self.test_db_path + extra
            if os.path.exists(p):
                try:
                    os.remove(p)
                except Exception:
                    pass

    async def test_t2_01_file_size_boundary_19_9mb_accepted(self):
        """[Boundary] 19.9MB payload must pass size check and not be rejected by size guard."""
        file_size_19_9mb = int(19.9 * 1024 * 1024)
        mock_doc = MockTelegramDocument("large_valid.txt", file_size=file_size_19_9mb, data=b"hello world")
        mock_update, mock_context = create_mock_update_message(document=mock_doc)

        with patch("bot.ProgressTracker.start", AsyncMock()), \
             patch("bot.ProgressTracker.finish", AsyncMock()) as mock_finish, \
             patch("bot.generate_chat_response", AsyncMock(return_value=ChatResponse(clean_text="تحلیل سند کامل شد."))), \
             patch("asyncio.sleep", AsyncMock()):

            await handle_document(mock_update, mock_context)

            # Assert size error was NOT thrown
            mock_finish.assert_called_once()
            finish_text = mock_finish.call_args[0][0]
            self.assertNotIn("بیش از ۲۰ مگابایت", finish_text)
            self.assertIn("تحلیل سند کامل شد.", finish_text)

    async def test_t2_02_file_size_boundary_20_1mb_rejected(self):
        """[Boundary] 20.1MB payload must be immediately rejected with Persian error notice."""
        file_size_20_1mb = int(20.1 * 1024 * 1024)
        mock_doc = MockTelegramDocument("oversized.pdf", file_size=file_size_20_1mb, data=b"huge data")
        mock_update, mock_context = create_mock_update_message(document=mock_doc)

        await handle_document(mock_update, mock_context)

        mock_update.message.reply_text.assert_called_once()
        reject_msg = mock_update.message.reply_text.call_args[0][0]
        self.assertIn("بیشتر از ۲۰ مگابایت است", reject_msg)
        self.assertIn("طبق قوانین تلگرام", reject_msg)

    async def test_t2_03_malformed_truncated_docx_payload(self):
        """[Boundary] Malformed or truncated DOCX bytes must be handled safely without uncaught exceptions."""
        truncated_bytes = b"PK\x03\x04\x14\x00\x00\x00\x08\x00corrupt_truncated_header_data"
        extracted = extract_docx_text(truncated_bytes)
        self.assertIsNone(extracted)

        # In handle_document, unreadable docx must return graceful notice
        mock_doc = MockTelegramDocument("corrupt.docx", file_size=len(truncated_bytes), data=truncated_bytes)
        mock_update, mock_context = create_mock_update_message(document=mock_doc)

        with patch("bot.ProgressTracker.start", AsyncMock()), \
             patch("bot.ProgressTracker.finish", AsyncMock()) as mock_finish, \
             patch("asyncio.sleep", AsyncMock()):

            await handle_document(mock_update, mock_context)

            mock_finish.assert_called_once()
            res_msg = mock_finish.call_args[0][0]
            self.assertIn("فایل Word باز شد اما متنی داخل آن پیدا نشد", res_msg)

    async def test_t2_04_malformed_corrupt_audio_payload(self):
        """[Boundary] Corrupted audio payload must be handled gracefully without crashing polling worker."""
        corrupt_audio_bytes = b"\x00\x00\x00\x18ftypmp42corrupted_bytes"
        mock_audio = MockTelegramAudio("corrupt.mp3", file_size=len(corrupt_audio_bytes), is_voice=False, data=corrupt_audio_bytes)
        mock_update, mock_context = create_mock_update_message(audio=mock_audio)

        with patch("bot.ProgressTracker.start", AsyncMock()), \
             patch("bot.ProgressTracker.finish", AsyncMock()) as mock_finish, \
             patch("bot.generate_multimodal_response", AsyncMock(side_effect=Exception("Invalid audio frame"))), \
             patch("asyncio.sleep", AsyncMock()):

            await handle_audio(mock_update, mock_context)

            mock_finish.assert_called_once()
            res_msg = mock_finish.call_args[0][0]
            self.assertIn("خطا در شنیدن یا تحلیل فایل صوتی", res_msg)

    async def test_t2_05_url_boundary_404_not_found(self):
        """[Boundary] 404 Not Found URL must return friendly Persian error without crashing."""
        mock_404_resp = MagicMock(status_code=404)
        with patch("httpx.AsyncClient.get", AsyncMock(return_value=mock_404_resp)):
            page_info, err = await fetch_web_page_content("https://example.com/nonexistent-article")
            self.assertIsNone(page_info)
            self.assertIsNotNone(err)
            self.assertIn("۴۰۴", err)
            self.assertIn("یافت نشد", err)

    async def test_t2_06_url_boundary_500_server_error(self):
        """[Boundary] 500 Internal Server Error must return friendly Persian error."""
        mock_500_resp = MagicMock(status_code=500)
        with patch("httpx.AsyncClient.get", AsyncMock(return_value=mock_500_resp)):
            page_info, err = await fetch_web_page_content("https://broken-server.internal/page")
            self.assertIsNone(page_info)
            self.assertIsNotNone(err)
            self.assertIn("500", normalize_digits(err))
            self.assertIn("خطا", err)

    async def test_t2_07_url_boundary_binary_stream_rejection(self):
        """[Boundary] Binary file stream URLs (e.g. .zip, .bin, octet-stream) must be rejected prior to HTML parsing."""
        mock_bin_resp = MagicMock(
            status_code=200,
            headers={"content-type": "application/octet-stream"}
        )
        with patch("httpx.AsyncClient.get", AsyncMock(return_value=mock_bin_resp)):
            page_info, err = await fetch_web_page_content("https://cdn.example.com/installer.bin")
            self.assertIsNone(page_info)
            self.assertIn("فایل باینری یا چندرسانه‌ای", err)

    async def test_t2_08_conversational_intent_isolation_no_reminder_hijack(self):
        """[Boundary] General conversational messages with numbers and hours must NOT be hijacked as reminders."""
        general_conversations = [
            "در سال ۲۰۲۶ ساعت ۴ بعدازظهر چه کنفرانسی برگزار شد؟",
            "تفاوت ساعت ۱۰ صبح و ۲ بعدازظهر در تابش خورشید چیست؟",
            "من ۳ ساعت درس خواندم و تست زدم.",
            "چرا قطار ساعت ۸ صبح حرکت نکرد؟",
        ]

        for prompt in general_conversations:
            mock_update, mock_context = create_mock_update_message(text=prompt)

            with patch("bot.ProgressTracker.start", AsyncMock()), \
                 patch("bot.ProgressTracker.finish", AsyncMock()) as mock_finish, \
                 patch("bot.generate_chat_response", AsyncMock(return_value=ChatResponse(clean_text="پاسخ هوش مصنوعی"))), \
                 patch("asyncio.sleep", AsyncMock()):

                await handle_message(mock_update, mock_context)

                # Ensure reminder confirmation was NOT sent
                if mock_update.message.reply_text.called:
                    call_text = mock_update.message.reply_text.call_args[0][0]
                    self.assertNotIn("یادآور با موفقیت در سیستم ثبت شد", call_text)

    async def test_t2_09_midnight_clock_boundary_mapping(self):
        """[Boundary] Midnight clock boundary ('۱۲ شب' / '۱۲ بامداد') must map strictly to 00:00 (midnight)."""
        base_dt = datetime.datetime(2026, 10, 16, 10, 0, 0, tzinfo=TEHRAN_TZ)  # 10:00 AM

        # Case 1: ۱۲ شب
        res_1 = parse_reminder_input_tier1("۱۲ شب یادم بنداز قرص بخورم", base_dt)
        self.assertIsNotNone(res_1)
        due_utc_1, text_1 = res_1
        # 00:00 Tehran (+03:30) of next day = 20:30 UTC of today
        self.assertEqual(due_utc_1, "2026-10-16T20:30:00Z")
        self.assertEqual(text_1, "قرص بخورم")

        # Case 2: ۱۲ بامداد
        res_2 = parse_reminder_input_tier1("۱۲ بامداد ریستارت سرور", base_dt)
        self.assertIsNotNone(res_2)
        due_utc_2, text_2 = res_2
        self.assertEqual(due_utc_2, "2026-10-16T20:30:00Z")
        self.assertEqual(text_2, "ریستارت سرور")

    async def test_t2_10_missing_fonts_graceful_fallback(self):
        """[Boundary] Missing TrueType font files must gracefully fall back without throwing unhandled exceptions."""
        with patch("os.path.exists", return_value=False):
            font_path = get_persian_font_path()
            self.assertIsNone(font_path)

            # Setup fonts fallback test
            setup_pdf_fonts()

            # Reshaping with no fonts must still return text without crashing
            reshaped = reshape_persian_text("متن تستی بدون فونت")
            self.assertIsNotNone(reshaped)


# =============================================================================
# Tier 3: Concurrency & Stress Testing
# =============================================================================

class Tier3ConcurrencyStressTests(unittest.TestCase):
    """
    Tier 3: Concurrency & Stress Testing
    Verifies multi-threaded database stress with 50 concurrent writer threads, zero lock errors, and rapid CRUD.
    """
    tier = 3

    def setUp(self):
        self.test_db_path = f"test_tier3_{uuid.uuid4().hex[:8]}.db"
        self.db = Database(self.test_db_path)
        bot.db = self.db

    def tearDown(self):
        for extra in ["", "-wal", "-shm"]:
            p = self.test_db_path + extra
            if os.path.exists(p):
                try:
                    os.remove(p)
                except Exception:
                    pass

    def test_t3_01_50_concurrent_sqlite_writers_zero_locks(self):
        """[Stress] 50 concurrent SQLite writer threads must execute simultaneous writes and reads with 0 database lock errors."""
        num_threads = 50
        ops_per_thread = 5
        errors = []
        lock_errors = []

        def writer_worker(tid: int):
            try:
                for i in range(ops_per_thread):
                    due_iso = f"2026-10-17T{tid % 24:02d}:{i % 60:02d}:00Z"
                    rem_id = self.db.add_reminder(
                        user_id=1000 + tid,
                        chat_id=2000 + tid,
                        message_id=i,
                        reminder_text=f"Stress Reminder from thread {tid} op {i}",
                        due_at_utc=due_iso
                    )
                    # Immediate read-back
                    due_list = self.db.get_due_reminders("2026-10-18T00:00:00Z")
                    # Update status
                    self.db.update_reminder_status(rem_id, "active")
            except sqlite3.OperationalError as oe:
                if "locked" in str(oe).lower():
                    lock_errors.append(oe)
                errors.append(oe)
            except Exception as e:
                errors.append(e)

        threads = [threading.Thread(target=writer_worker, args=(t,)) for t in range(num_threads)]
        for th in threads:
            th.start()
        for th in threads:
            th.join(timeout=20.0)

        self.assertEqual(len(lock_errors), 0, f"Encountered database lock errors: {lock_errors}")
        self.assertEqual(len(errors), 0, f"Encountered general concurrency errors: {errors}")

        # Verify exact record count
        total_rows = self.db.execute("SELECT COUNT(*) as cnt FROM reminders")
        self.assertEqual(total_rows[0]["cnt"], num_threads * ops_per_thread)

    def test_t3_02_rapid_fire_reminder_crud_stress(self):
        """[Stress] Rapid-fire reminder scheduling, indexing queries, and status updates under heavy sequential burst."""
        total_items = 100
        for i in range(total_items):
            rem_id = self.db.add_reminder(
                user_id=i,
                chat_id=10000 + i,
                message_id=i,
                reminder_text=f"Rapid Reminder {i}",
                due_at_utc="2026-10-16T12:00:00Z"
            )
            self.assertTrue(rem_id > 0)

        # Batch query matching index
        due_reminders = self.db.get_due_reminders("2026-10-16T13:00:00Z", limit=total_items)
        self.assertEqual(len(due_reminders), total_items)

        # Bulk status transition
        for r in due_reminders:
            self.db.update_reminder_status(r["id"], "delivered")

        # Query again: must be empty
        rem_after = self.db.get_due_reminders("2026-10-16T13:00:00Z", limit=total_items)
        self.assertEqual(len(rem_after), 0)

    def test_t3_03_scheduler_worker_polling_under_heavy_concurrency(self):
        """[Stress] Scheduler worker single pass must execute safely while concurrent background threads insert reminders."""
        stop_event = threading.Event()
        writer_errors = []

        def background_inserter():
            count = 0
            while not stop_event.is_set() and count < 30:
                try:
                    self.db.add_reminder(
                        user_id=999,
                        chat_id=888,
                        message_id=count,
                        reminder_text=f"Background async reminder {count}",
                        due_at_utc="2026-10-16T10:00:00Z"
                    )
                    count += 1
                    time.sleep(0.005)
                except Exception as e:
                    writer_errors.append(e)

        bg_thread = threading.Thread(target=background_inserter)
        bg_thread.start()

        mock_app = MagicMock()
        mock_app.bot.send_message = AsyncMock()

        try:
            # Run scheduler loop single pass
            asyncio.run(reminder_scheduler_loop(mock_app, single_pass=True))
        finally:
            stop_event.set()
            bg_thread.join(timeout=5.0)

        self.assertEqual(len(writer_errors), 0, f"Background inserter had errors: {writer_errors}")


# =============================================================================
# Tier 4: Resilience & Fault Tolerance
# =============================================================================

class Tier4ResilienceFaultToleranceTests(unittest.IsolatedAsyncioTestCase):
    """
    Tier 4: Resilience & Fault Tolerance
    Verifies 503 fallback, 429 backoff, server reboot recovery (<24h vs >24h), and Markdown parse fallback.
    """
    tier = 4

    def setUp(self):
        self.test_db_path = f"test_tier4_{uuid.uuid4().hex[:8]}.db"
        self.db = Database(self.test_db_path)
        bot.db = self.db
        self.user_id = "14004"
        self.db.get_or_create_user(self.user_id)
        self.active_chat = self.db.get_active_chat(self.user_id)

    def tearDown(self):
        for extra in ["", "-wal", "-shm"]:
            p = self.test_db_path + extra
            if os.path.exists(p):
                try:
                    os.remove(p)
                except Exception:
                    pass

    async def test_t4_01_gemini_503_transient_spike_fallback_recovery(self):
        """[Resilience] Gemini 503 UNAVAILABLE transient spike must recover automatically via fallback models."""
        attempt_counter = [0]

        def fake_build_gemini_session(user_id, chat_id, target_model, is_fallback=False, enable_code_exec=True):
            attempt_counter[0] += 1
            mock_session = MagicMock()
            if attempt_counter[0] == 1:
                # Primary model throws 503
                mock_session.send_message.side_effect = Exception("503 UNAVAILABLE: Model overloaded")
            else:
                # Fallback model succeeds
                mock_session.send_message.return_value = MockGenAIResponse(
                    parts=[MockGenAIPart(text="پاسخ تایید شده از مدل جایگزین")],
                    text="پاسخ تایید شده از مدل جایگزین"
                )
            return mock_session

        with patch("bot.build_gemini_session", side_effect=fake_build_gemini_session), \
             patch("asyncio.sleep", AsyncMock()):

            resp = await generate_chat_response(self.user_id, self.active_chat["id"], "سوال تست")
            self.assertGreaterEqual(attempt_counter[0], 2)
            self.assertIn("مدل جایگزین", resp.clean_text)

    async def test_t4_02_gemini_429_rate_limit_backoff_recovery(self):
        """[Resilience] Gemini 429 RESOURCE_EXHAUSTED rate limit must trigger backoff and recover."""
        attempt_counter = [0]

        def fake_build_session(user_id, chat_id, target_model, is_fallback=False, enable_code_exec=True):
            attempt_counter[0] += 1
            mock_session = MagicMock()
            if attempt_counter[0] == 1:
                mock_session.send_message.side_effect = Exception("429 RESOURCE_EXHAUSTED: Quota exceeded")
            else:
                mock_session.send_message.return_value = MockGenAIResponse(
                    parts=[MockGenAIPart(text="پاسخ موفق پس از بک‌آف")],
                    text="پاسخ موفق پس از بک‌آف"
                )
            return mock_session

        with patch("bot.build_gemini_session", side_effect=fake_build_session), \
             patch("asyncio.sleep", AsyncMock()) as mock_sleep:

            resp = await generate_chat_response(self.user_id, self.active_chat["id"], "سوال پس از لیمیت")
            self.assertGreaterEqual(attempt_counter[0], 2)
            mock_sleep.assert_called()
            self.assertIn("پاسخ موفق پس از بک‌آف", resp.clean_text)

    async def test_t4_03_server_reboot_crash_recovery_overdue_reminders(self):
        """[Resilience] Reboot crash recovery: overdue <24h delivered with warning banner; overdue >24h marked missed."""
        now_dt = datetime.datetime.now(datetime.timezone.utc)

        # 1. Reminders overdue by 2 hours (<24h) -> Must be delivered with delay banner
        due_recent = (now_dt - datetime.timedelta(hours=2)).strftime("%Y-%m-%dT%H:%M:%SZ")
        id_recent = self.db.add_reminder(
            user_id=1401,
            chat_id=2401,
            message_id=10,
            reminder_text="جلسه اضطراری",
            due_at_utc=due_recent
        )

        # 2. Reminders overdue by 30 hours (>24h) -> Must be marked missed
        due_old = (now_dt - datetime.timedelta(hours=30)).strftime("%Y-%m-%dT%H:%M:%SZ")
        id_old = self.db.add_reminder(
            user_id=1402,
            chat_id=2402,
            message_id=20,
            reminder_text="خرید نان دیروز",
            due_at_utc=due_old
        )

        mock_app = MagicMock()
        mock_app.bot.send_message = AsyncMock()

        await recover_server_reboot_reminders(mock_app)

        # Check DB states
        rows = self.db.execute("SELECT id, status FROM reminders WHERE id IN (?, ?)", (id_recent, id_old))
        status_map = {r["id"]: r["status"] for r in rows}

        self.assertEqual(status_map[id_recent], "delivered")
        self.assertEqual(status_map[id_old], "missed")

        # Verify delivery text included warning banner
        mock_app.bot.send_message.assert_called_once()
        sent_text = mock_app.bot.send_message.call_args[1]["text"]
        self.assertIn("یادآوری با تأخیر", sent_text)
        self.assertIn("جلسه اضطراری", sent_text)

    async def test_t4_04_scheduler_markdown_entity_parse_error_resilience(self):
        """[Resilience] Scheduler loop entity parsing resilience: BadRequest on Markdown gracefully falls back to plain text."""
        # Insert a due reminder
        now_utc = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        rem_id = self.db.add_reminder(
            user_id=1403,
            chat_id=2403,
            message_id=15,
            reminder_text="متن با کاراکترهای خراب *ناقص _bold",
            due_at_utc=now_utc
        )

        mock_app = MagicMock()
        first_attempt = [True]

        async def fake_send_message(chat_id, text, parse_mode=None, reply_markup=None, reply_to_message_id=None):
            if first_attempt[0] and parse_mode == "Markdown":
                first_attempt[0] = False
                raise BadRequest("Can't parse entities: can't find end of bold entity")
            return MagicMock()

        mock_app.bot.send_message = AsyncMock(side_effect=fake_send_message)

        await reminder_scheduler_loop(mock_app, single_pass=True)

        # Verify fallback to plain text occurred and reminder is delivered
        rows = self.db.execute("SELECT status FROM reminders WHERE id = ?", (rem_id,))
        self.assertEqual(rows[0]["status"], "delivered")
        self.assertFalse(first_attempt[0])


# =============================================================================
# Tier 5: Real-World Workload Scenarios
# =============================================================================

class Tier5RealWorldWorkloadTests(unittest.IsolatedAsyncioTestCase):
    """
    Tier 5: Real-World Workload Scenarios
    Full end-to-end multi-turn conversation flow exercising code execution, charting, exports, and reminders.
    """
    tier = 5

    def setUp(self):
        self.test_db_path = f"test_tier5_{uuid.uuid4().hex[:8]}.db"
        self.db = Database(self.test_db_path)
        bot.db = self.db
        self.user_id = "15005"
        self.db.get_or_create_user(self.user_id)
        self.active_chat = self.db.get_active_chat(self.user_id)

    def tearDown(self):
        for extra in ["", "-wal", "-shm"]:
            p = self.test_db_path + extra
            if os.path.exists(p):
                try:
                    os.remove(p)
                except Exception:
                    pass

    async def test_t5_01_e2e_computational_and_data_export_flow(self):
        """[Workload] End-to-end flow: Computational Python solving followed by dynamic Excel and Chart generation."""
        # Turn 1: Computational query
        math_prompt = "میانگین و انحراف معیار اعداد [10, 20, 30, 40, 50] را با پایتون حساب کن."
        mock_update_1, mock_context_1 = create_mock_update_message(text=math_prompt)

        math_code = "import statistics\ndata = [10, 20, 30, 40, 50]\nprint(f'Mean: {statistics.mean(data)}, Stdev: {statistics.stdev(data):.2f}')"
        math_result = "Mean: 30, Stdev: 15.81\n"
        math_answer = "میانگین برابر ۳۰ و انحراف معیار برابر ۱۵.۸۱ است."

        mock_parts = [
            MockGenAIPart(thought=True, text="محاسبه آماری"),
            MockGenAIPart(executable_code=math_code),
            MockGenAIPart(code_execution_result=math_result),
            MockGenAIPart(text=math_answer),
        ]
        mock_math_resp = MockGenAIResponse(parts=mock_parts, text=math_answer)

        with patch("bot.ProgressTracker.start", AsyncMock()), \
             patch("bot.ProgressTracker.finish", AsyncMock()) as mock_finish_1, \
             patch("bot.build_gemini_session") as mock_build_session, \
             patch("asyncio.sleep", AsyncMock()):

            mock_session = MagicMock()
            mock_session.send_message.return_value = mock_math_resp
            mock_build_session.return_value = mock_session

            await handle_message(mock_update_1, mock_context_1)

            mock_finish_1.assert_called_once()
            final_turn1_text = mock_finish_1.call_args[0][0]
            self.assertIn("💻 **کد پایتون اجرا شده:**", final_turn1_text)
            self.assertIn("statistics.mean(data)", final_turn1_text)
            self.assertIn("Stdev: 15.81", final_turn1_text)

        # Turn 2: Export chat history into Excel and Chart
        export_dataset = [
            {"شاخص": "میانگین", "مقدار": 30},
            {"شاخص": "انحراف معیار", "مقدار": 15.81},
        ]
        excel_file = generate_excel_report(export_dataset, title="گزارش_آماری")
        self.assertIsInstance(excel_file, io.BytesIO)
        self.assertTrue(excel_file.getvalue().startswith(b"PK\x03\x04"))

        chart_data = {"labels": ["میانگین", "انحراف معیار"], "values": [30, 15.81]}
        chart_file = generate_chart(chart_data, chart_type="bar")
        self.assertIsInstance(chart_file, io.BytesIO)
        self.assertTrue(chart_file.getvalue().startswith(b"\x89PNG"))

    async def test_t5_02_e2e_bilingual_reminder_and_reasoning_flow(self):
        """[Workload] End-to-end flow: Bilingual reminder scheduling, persistence, and transparent reasoning inspection."""
        # Turn 1: User schedules persistent reminder
        user_prompt = "فردا ساعت ۸ صبح جلسه هماهنگی تیم رو یادم بنداز"
        mock_update, mock_context = create_mock_update_message(user_id=int(self.user_id), text=user_prompt)

        with patch("bot.ProgressTracker.start", AsyncMock()), \
             patch("bot.ProgressTracker.finish", AsyncMock()), \
             patch("asyncio.sleep", AsyncMock()):

            await handle_message(mock_update, mock_context)

            # Assert reminder confirmation was returned
            mock_update.message.reply_text.assert_called_once()
            reply_msg = mock_update.message.reply_text.call_args[0][0]
            self.assertIn("یادآور با موفقیت در سیستم ثبت شد", reply_msg)
            self.assertIn("جلسه هماهنگی تیم", reply_msg)

        # Verify persistence in SQLite
        due_rems = self.db.execute("SELECT * FROM reminders WHERE user_id = ?", (int(self.user_id),))
        self.assertEqual(len(due_rems), 1)
        self.assertEqual(due_rems[0]["reminder_text"], "جلسه هماهنگی تیم")
        self.assertEqual(due_rems[0]["status"], "pending")

        # Turn 2: Reasoning log retrieval and inspection
        thought_id = "e2e_thought_123"
        self.db.save_thought(
            thought_id=thought_id,
            user_id=self.user_id,
            chat_id=self.active_chat["id"],
            thought_content="تحلیل نیازهای پروژه و تعیین اولویت‌های تسک‌ها برای جلسه فردا."
        )

        mock_cb_query = MagicMock(spec=CallbackQuery)
        mock_cb_query.data = f"show_thought:{thought_id}"
        mock_cb_query.message = MagicMock(spec=Message)
        mock_cb_query.message.reply_text = AsyncMock()
        mock_cb_query.answer = AsyncMock()

        mock_cb_update = MagicMock(spec=Update)
        mock_cb_update.callback_query = mock_cb_query
        mock_cb_context = MagicMock()

        await callback_handler(mock_cb_update, mock_cb_context)

        mock_cb_query.message.reply_text.assert_called_once()
        thought_out = mock_cb_query.message.reply_text.call_args[0][0]
        self.assertIn("تحلیل نیازهای پروژه", thought_out)


# =============================================================================
# Custom Colored Test Runner & Scorecard Formatter
# =============================================================================

class ANSIColor:
    RESET = "\033[0m"
    BOLD = "\033[1m"
    GREEN = "\033[92m"
    RED = "\033[91m"
    YELLOW = "\033[93m"
    BLUE = "\033[94m"
    CYAN = "\033[96m"
    WHITE = "\033[97m"


class AdversarialScorecardTestResult(unittest.TextTestResult):
    """Custom TestResult that captures timing, status, and tier stats for scorecard generation."""

    TIER_NAMES = {
        1: "Feature Isolation & Unit Verification",
        2: "Boundary & Corner Cases",
        3: "Concurrency & Stress Testing",
        4: "Resilience & Fault Tolerance",
        5: "Real-World Workload Scenarios",
    }

    def __init__(self, stream, descriptions, verbosity):
        super().__init__(stream, descriptions, verbosity)
        self.tier_stats = {
            t: {"name": self.TIER_NAMES[t], "run": 0, "passed": 0, "failed": 0, "time": 0.0}
            for t in range(1, 6)
        }
        self.test_start_time = 0.0

    def startTest(self, test):
        super().startTest(test)
        self.test_start_time = time.perf_counter()

    def _get_tier(self, test) -> int:
        test_class = getattr(test, "__class__", None)
        return getattr(test_class, "tier", 1)

    def addSuccess(self, test):
        elapsed = time.perf_counter() - self.test_start_time
        super().addSuccess(test)
        tier = self._get_tier(test)
        self.tier_stats[tier]["run"] += 1
        self.tier_stats[tier]["passed"] += 1
        self.tier_stats[tier]["time"] += elapsed

        doc = test.shortDescription() or test._testMethodName
        print(f"  {ANSIColor.GREEN}[PASS]{ANSIColor.RESET} Tier {tier}: {doc} ({elapsed:.3f}s)")

    def addFailure(self, test, err):
        elapsed = time.perf_counter() - self.test_start_time
        super().addFailure(test, err)
        tier = self._get_tier(test)
        self.tier_stats[tier]["run"] += 1
        self.tier_stats[tier]["failed"] += 1
        self.tier_stats[tier]["time"] += elapsed

        doc = test.shortDescription() or test._testMethodName
        print(f"  {ANSIColor.RED}[FAIL]{ANSIColor.RESET} Tier {tier}: {doc} ({elapsed:.3f}s)")

    def addError(self, test, err):
        elapsed = time.perf_counter() - self.test_start_time
        super().addError(test, err)
        tier = self._get_tier(test)
        self.tier_stats[tier]["run"] += 1
        self.tier_stats[tier]["failed"] += 1
        self.tier_stats[tier]["time"] += elapsed

        doc = test.shortDescription() or test._testMethodName
        print(f"  {ANSIColor.RED}[ERROR]{ANSIColor.RESET} Tier {tier}: {doc} ({elapsed:.3f}s)")

    def print_scorecard(self):
        total_run = sum(s["run"] for s in self.tier_stats.values())
        total_passed = sum(s["passed"] for s in self.tier_stats.values())
        total_failed = sum(s["failed"] for s in self.tier_stats.values())
        total_time = sum(s["time"] for s in self.tier_stats.values())

        print("\n" + "=" * 82)
        print(f"{ANSIColor.BOLD}{ANSIColor.CYAN}            ADVERSARIAL STRESS QA SUITE SCORECARD (REQUIREMENT R7){ANSIColor.RESET}")
        print("=" * 82)
        print(f" {'Tier':<6} {'Tier Name':<42} {'Tests':<8} {'Passed':<8} {'Failed':<8} {'Time (s)':<10}")
        print("-" * 82)

        for t in range(1, 6):
            s = self.tier_stats[t]
            color = ANSIColor.GREEN if s["failed"] == 0 and s["run"] > 0 else ANSIColor.RED
            tier_code = f"T{t}"
            print(f" {tier_code:<6} {s['name']:<42} {s['run']:<8} {color}{s['passed']:<8}{ANSIColor.RESET} {s['failed']:<8} {s['time']:<10.2f}s")

        print("-" * 82)
        status_color = ANSIColor.GREEN if total_failed == 0 and total_run > 0 else ANSIColor.RED
        print(f" {'TOTAL':<6} {'All 5 Tiers Combined':<42} {total_run:<8} {status_color}{total_passed:<8}{ANSIColor.RESET} {total_failed:<8} {total_time:<10.2f}s")
        print("=" * 82)

        if total_failed == 0 and total_run > 0:
            print(f" {ANSIColor.BOLD}{ANSIColor.GREEN}OVERALL STATUS: ALL TIERS PASSED (100% SUCCESS - READY FOR AUDIT){ANSIColor.RESET}")
        else:
            print(f" {ANSIColor.BOLD}{ANSIColor.RED}OVERALL STATUS: FAILURE DETECTED IN QA SUITE ({total_failed} FAILED){ANSIColor.RESET}")
        print("=" * 82 + "\n")


class AdversarialScorecardRunner:
    """Custom TestRunner executing all 5 tiers and producing the scorecard."""
    def run(self, suite):
        result = AdversarialScorecardTestResult(sys.stdout, True, 1)
        suite.run(result)
        result.print_scorecard()
        return result


def build_full_suite() -> unittest.TestSuite:
    suite = unittest.TestSuite()
    loader = unittest.TestLoader()
    suite.addTests(loader.loadTestsFromTestCase(Tier1FeatureIsolationTests))
    suite.addTests(loader.loadTestsFromTestCase(Tier2BoundaryCornerCaseTests))
    suite.addTests(loader.loadTestsFromTestCase(Tier3ConcurrencyStressTests))
    suite.addTests(loader.loadTestsFromTestCase(Tier4ResilienceFaultToleranceTests))
    suite.addTests(loader.loadTestsFromTestCase(Tier5RealWorldWorkloadTests))
    return suite


if __name__ == "__main__":
    print(f"\n{ANSIColor.BOLD}{ANSIColor.BLUE}Starting Comprehensive Adversarial QA Suite (Requirement R7)...{ANSIColor.RESET}\n")
    full_suite = build_full_suite()
    runner = AdversarialScorecardRunner()
    test_result = runner.run(full_suite)
    sys.exit(0 if test_result.wasSuccessful() else 1)
